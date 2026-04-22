import re
import os
import random
import pytest
from playwright.sync_api import Page, expect
from datetime import datetime, timedelta
from openpyxl import load_workbook, Workbook

RESULTS_FILE = os.path.join(os.path.dirname(__file__), "..", "..", "..", "results", "fg_test_results.xlsx")

def _append_to_xlsx(row: list):
    path = os.path.abspath(RESULTS_FILE)
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["User Name", "Address", "Product Type", "Coverage A Limit", "Quote Number", "Premium"])
    ws.append(row)
    wb.save(path)


FL_ADDRESSES = [
    "8734 Crestgate Cir, Orlando, FL 32819",
    "19240 Pepper Grass Dr, Tampa, FL 33647",
    "8832 Savona Ct, Naples, FL 34119",
    "1461 75th Cir NE, Saint Petersburg, FL 33702",
    "8050 Princeton Dr, Naples, FL 34104",
    "12274 Accipiter Dr, Orlando, FL 32837",
    "15110 Shearcrest Dr, Lithia, FL 33547",
    "7733 Belvoir Dr, Orlando, FL 32835",
    "9426 Meadow Crest Ln, Clermont, FL 34711",
    "612 Citrus Wood Ln, Valrico, FL 33594",
    "66 Woodbury Dr, Palm Coast, FL 32164",
    "1401 Maine Ave, Lynn Haven, FL 32444",
    "9607 Bahia Rd, Ocala, FL 34472",
    "2150 NW 74th Ave, Fort Lauderdale, FL 33313",
    "4364 Tarragon Ave, Middleburg, FL 32068",
    "3198 Michaels Ct, Green Cove Springs, FL 32043",
    "3250 Dowitcher Ln, Orange Park, FL 32065",
    "613 S West St, Green Cove Springs, FL 32043",
    "8329 Lakemont Dr, Jacksonville, FL 32216",
    "13501 Princess Kelly Dr, Jacksonville, FL 32225",
    "6229 Cranberry Ln E, Jacksonville, FL 32244",
    "14688 Indigo Lakes Cir, Naples, FL 34119",
]

def test_example(page: Page, snap) -> None:
    username = "icortesmarquez@geovera"
    address = random.choice(FL_ADDRESSES)
    coverage_a = "800,000"

    page.goto("https://tst.mygeosource.com/")
    page.get_by_role("textbox", name="Policyholder Email or Agent").click()
    page.get_by_role("textbox", name="Policyholder Email or Agent").fill(username)
    page.get_by_role("textbox", name="Password").click()
    page.get_by_role("textbox", name="Password").fill(os.environ["INTERNAL_PASSWORD"])
    page.get_by_role("button", name="Log In", exact=True).click()
    snap(page, "01_login")
    page.get_by_role("paragraph").filter(has_text="Start New Quote").get_by_role("link").click()
    page.locator("#optionHomeowners").get_by_text("GeoVera Specialty Insurance").click()
    snap(page, "02_product_select")
    with page.expect_popup() as page1_info:
        page.get_by_role("button", name="OK").click()
    page1 = page1_info.value
    page1.wait_for_selector("#ctl00_WelcomeUserControl_ddlWholesaler", timeout=30000)
    page1.locator("#ctl00_WelcomeUserControl_ddlWholesaler").select_option("APPUND")
    page1.locator("#ctl00_WelcomeUserControl_txtProducerNumber").fill("APPAGT")
    page1.locator("#ctl00_WelcomeUserControl_btnUpdateQuotingAs").click()
    page1.wait_for_selector("#ctl00_WelcomeUserControl_ucWelcomeUserLookup_dgUserList_ctl02_btnId", timeout=30000)
    page1.locator("#ctl00_WelcomeUserControl_ucWelcomeUserLookup_dgUserList_ctl02_btnId").click()
    page1.wait_for_timeout(3000)
    snap(page1, "03_welcome")
    page1.locator("#RadAutoCompleteBoxGeo_Input").click()
    page1.locator("#RadAutoCompleteBoxGeo_Input").press_sequentially(address)
    page1.wait_for_timeout(1000)
    page1.locator("#RadAutoCompleteBoxGeo_Input").press("Enter")
    snap(page1, "04_address")
    page1.get_by_role("button", name="Next -->").click()
    page1.get_by_role("checkbox", name="Auros HO-3").uncheck()
    page1.get_by_role("checkbox", name="HO Wind Only").uncheck()
    page1.locator("#txtRequestedHomeownersCoverageALimit").click()
    page1.locator("#txtRequestedHomeownersCoverageALimit").fill(coverage_a)
    snap(page1, "05_coverages")
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtYearBuilt").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtYearBuilt").fill("2025")
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtLivingArea").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtLivingArea").fill("2025")
    page1.get_by_role("button", name="Next -->").click()
    page1.locator("#ddlConstructionType").select_option("FR")
    page1.locator("#ddlSidingTypeDominant").select_option("ST")
    page1.locator("#ddlRestrictedRoofShapes").select_option("False")
    if page1.locator("#ddlRoofShapeDominant").is_visible():
        page1.locator("#ddlRoofShapeDominant").select_option("2")
    snap(page1, "06_property")
    page1.get_by_role("button", name="Next -->").click()
    page1.wait_for_timeout(3000)
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlCentralFireAlarm").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlCentralBurglarAlarm").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlGatedGuardedCommunity").select_option("False")
    if page1.locator("#ddlSwimmingPool").is_visible():
        page1.locator("#ddlSwimmingPool").select_option(label="No Swimming Pool")
    snap(page1, "07_protection")
    page1.get_by_role("button", name="Next -->").click()
    page1.locator("#ddlPolybutylenePipes").select_option("False")
    page1.locator("#ddlAutomaticWaterShutoffValve").select_option("False")
    snap(page1, "08_plumbing")
    page1.get_by_role("button", name="Next -->").click()
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlWindMitigationVerificationInspection").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlWindMitigationVerificationInspection").select_option("Yes")
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtWindMitigationVerificationInspectionMonth").fill("01")
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtWindMitigationVerificationInspectionDay").fill("01")
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtWindMitigationVerificationInspectionYear").fill("2022")
        page1.locator("#fileWindMitigationVerificationInspectionForm").set_input_files("C:/Users/icortesmarquez/Desktop/Playwright/HO_PW/test_data/testfile.pdf")
    snap(page1, "09_wind_mitigation")
    page1.get_by_role("button", name="Next -->").click()
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtFirstName").click()
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtFirstName").press("CapsLock")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtFirstName").press_sequentially("ICMTEST")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtLastName").click()
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtLastName").press_sequentially("TESTER")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtDOBMonth").click()
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtDOBMonth").press_sequentially("01")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtDOBDay").press_sequentially("01")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtDOBYear").press_sequentially("1998")
    page1.locator("input[name=\"ctl00\\$ContentPlaceHolder1\\$ctl05\\$txtSSN1\"]").press_sequentially("123")
    page1.locator("input[name=\"ctl00\\$ContentPlaceHolder1\\$ctl05\\$txtSSN2\"]").press_sequentially("11")
    page1.locator("input[name=\"ctl00\\$ContentPlaceHolder1\\$ctl05\\$txtSSN3\"]").press_sequentially("1122")
    page1.get_by_role("checkbox", name="Is Mailing Address the same").check()
    page1.locator("#ddlPriorHomeownersInsurance").select_option("False")
    page1.wait_for_timeout(3000)
    page1.locator("#ddlPriorHomeownersInsurance").select_option("False")
    snap(page1, "10_applicant")
    page1.get_by_role("button", name="Next -->").click()
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl00_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl01_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl02_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl03_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl04_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl05_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl06_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl07_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl08_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl09_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl10_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl11_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl12_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl13_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl14_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl15_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl16_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl17_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl18_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl19_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl20_ddlQuestions").select_option("False")
    page1.get_by_role("cell", name="Applicant acknowledges that flood coverage is excluded for the homeowners, wind-only, and vacant policy.", exact=True).locator("span").first.click()
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkRoofPaymentSchedule").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkRoofPaymentSchedule").check()
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkFloodExclusion").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkFloodExclusion").check()
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkRoofExclusion").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkRoofExclusion").check()
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkAcknowledgedWaterDamageLimitation").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkAcknowledgedWaterDamageLimitation").check()
    snap(page1, "11_questions")
    page1.get_by_role("button", name="Submit Quote Request").click()


    geovera_panel = page1.locator("div[id$='_pnlHOQuotes']").filter(has_text="GeoVera Specialty Insurance Company")
    geovera_panel.wait_for(state="visible", timeout=30000)
    snap(page1, "12_quote_results")
    ProductName = geovera_panel.locator("span[id$='_lblQuote1Header']").inner_text()
    print(f'Homeowners Product Option: {ProductName}')
    QuoteNumber = geovera_panel.locator("span[id$='_lblQuote1Number']").inner_text()
    print(f'Quote Number: {QuoteNumber}')
    FullPay = geovera_panel.locator("span[id$='_lstPaymentOptions_ctl00_lblQuote1PaymentOption']").inner_text()
    print(f'Full Pay Option: {FullPay}')
    premium_match = re.search(r"\$\d{1,3}(?:,\d{3})*(?:\.\d{2})", FullPay)
    premium = premium_match.group() if premium_match else FullPay

    ThreePayFirstInstallment = geovera_panel.locator(
        "span[id$='_lstPaymentOptions_ctl01_lblQuote1PaymentOption']"
    ).text_content().strip()

    match = re.search(r"\$\d{1,3}(?:,\d{3})*(?:\.\d{2})", ThreePayFirstInstallment)

    if match:
        print(f"3 Pay First Installment:\n{match.group().strip()}\n")
    else:
        print("No dollar amount found in:", ThreePayFirstInstallment)

    page1.get_by_role("button", name="Continue to Bind Coverage").click()
    page1.locator("#chkPopInConfirmation").check()
    page1.get_by_role("button", name="Continue", exact=True).click()

    eff_date = (datetime.today()).strftime('%m/%d/%Y')
    page1.locator("#ctl00_ContentPlaceHolder1_ucBinderRequest_txtBinderEffectiveDate1").fill(eff_date)
    page1.locator("#ctl00_ContentPlaceHolder1_ucBinderRequest_ddlBillTo1").select_option("0")
    page1.locator("#ddlPaymentPlan1").select_option("0")
    snap(page1, "13_bind_request")
    page1.get_by_role("button", name="Request to Bind").click()

    page1.wait_for_selector("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblAmountDue", timeout=30000)
    snap(page1, "14_binder_detail")
    BinderNumber = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblBinderNumber").inner_text()
    print(f'Binder Number: {BinderNumber}')
    BillTo = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblBillTo").inner_text()
    print(f'Bill To: {BillTo}')
    PaymentPlan = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblPaymentPlan").inner_text()
    print(f'Payment Plan: {PaymentPlan}')
    EffectiveDate = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblEffectiveDate").inner_text()
    print(f'Effective Date: {EffectiveDate}')
    AmountDueNow = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblAmountDue").inner_text()
    print(f'Amount Due Now: {AmountDueNow}')
    page1.wait_for_timeout(5000)

    _append_to_xlsx([username, address, ProductName, coverage_a, QuoteNumber, premium])
