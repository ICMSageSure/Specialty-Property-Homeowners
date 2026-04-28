import os
import random
from openpyxl import load_workbook, Workbook
from playwright.sync_api import Page

from pages.login_page import LoginPage
from pages.quote_launch_page import QuoteLaunchPage
from pages.select_wholesaler_page import SelectWholesalerPage
from pages.steps import (
    AddressStep,
    CoveragesStep,
    PropertyStep,
    ProtectionStep,
    PlumbingStep,
    WindMitigationStep,
    ApplicantStep,
    QuestionsStep,
)
from pages.quote_results_page import QuoteResultsPage, BindRequestPage, BinderDetailPage

RESULTS_FILE = os.path.join(os.path.dirname(__file__), "..", "..", "..", "results", "fg_test_results.xlsx")

FL_ADDRESSES = [
    "8734 Crestgate Cir, Orlando, FL 32819",
    "19240 Pepper Grass Dr, Tampa, FL 33647",
    "8832 Savona Ct, Naples, FL 34119",
    "1461 75th Cir NE, Saint Petersburg, FL 33702",
    "8050 Princeton Dr, Naples, FL 34104",
    "12274 Accipiter Dr, Orlando, FL 32837",
    "15110 Shearcrest Dr, Lithia, FL 33547",
    "7733 Belvoir Dr, Orlando, FL 32835",
    "9426 Meadow Crest Ln, Clermont, FL 34711",
    "612 Citrus Wood Ln, Valrico, FL 33594",
    "66 Woodbury Dr, Palm Coast, FL 32164",
    "1401 Maine Ave, Lynn Haven, FL 32444",
    "9607 Bahia Rd, Ocala, FL 34472",
    "2150 NW 74th Ave, Fort Lauderdale, FL 33313",
    "4364 Tarragon Ave, Middleburg, FL 32068",
    "3198 Michaels Ct, Green Cove Springs, FL 32043",
    "3250 Dowitcher Ln, Orange Park, FL 32065",
    "613 S West St, Green Cove Springs, FL 32043",
    "8329 Lakemont Dr, Jacksonville, FL 32216",
    "13501 Princess Kelly Dr, Jacksonville, FL 32225",
    "6229 Cranberry Ln E, Jacksonville, FL 32244",
    "14688 Indigo Lakes Cir, Naples, FL 34119",
]


def _append_to_xlsx(row: list):
    path = os.path.abspath(RESULTS_FILE)
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["User Name", "Address", "Test Scenario", "Product Type", "Coverage A Limit", "Quote Number", "Quote Expiration", "Dwelling (Cov A)", "Total Annual Premium", "Full Pay", "3 Pay (1st Install)", "4 Pay (1st Install)", "Monthly (1st Install)", "Name (max 30)", "Second Name (max 30)", "Street (max 30)", "City (max 30)", "Zip (max 5)", "Zip+4 (max 4)", "Binder Number"])
    ws.append(row)
    wb.save(path)


def test_quote_fg_ho_wind_v2(page: Page):
    username = "icortesmarquez@geovera.com"
    address = random.choice(FL_ADDRESSES)
    coverage_a = "800,000"

    LoginPage(page).login(username, os.environ["INTERNAL_PASSWORD"])

    popup = QuoteLaunchPage(page).start_ho_quote()

    SelectWholesalerPage(popup).setup_wholesaler()
    AddressStep(popup).enter_address(address)
    CoveragesStep(popup).select_wind_only(coverage_a)
    PropertyStep(popup).fill()
    ProtectionStep(popup).fill()
    PlumbingStep(popup).fill()
    WindMitigationStep(popup).fill()
    ApplicantStep(popup).fill()
    QuestionsStep(popup).answer_all_false_and_submit()

    results = QuoteResultsPage(popup)
    results.wait_for_results()
    results.print_quote_summary()
    product_name = results.get_product_name()
    quote_number = results.get_quote_number()
    premium = results.get_premium()
    full_pay = results.get_full_pay()
    three_pay = results.get_three_pay()
    results.continue_to_bind()

    BindRequestPage(popup).fill_and_submit()
    binder_detail = BinderDetailPage(popup)
    binder_detail.wait_and_print()
    binder_number = binder_detail.get_binder_number()

    _append_to_xlsx([username, address, "FG HO Wind", product_name, coverage_a, quote_number, "", "", premium, full_pay, three_pay, "", "", "", "", "", "", "", "", binder_number])
