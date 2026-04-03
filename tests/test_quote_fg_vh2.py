import re
import os
import random
import pytest
from playwright.sync_api import Page, expect
from datetime import datetime, timedelta
from openpyxl import load_workbook, Workbook

RESULTS_FILE = os.path.join(os.path.dirname(__file__), "..", "fg_test_results.xlsx")

def _append_to_xlsx(row: list):
    path = os.path.abspath(RESULTS_FILE)
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["User Name", "Address", "Product Type", "Coverage A Limit", "Quote Number", "Premium"])
    ws.append(row)
    wb.save(path)


FL_ADDRESSES = [
    "72 Saint Kitts Cir, Winter Haven, FL 33884",
    "320 Sea Breeze Dr, Panama City Beach, FL 32413",
    "15443 Perdido Dr, Orlando, FL 32828",
    "19240 Pepper Grass Dr, Tampa, FL 33647",
    "450 SW 125th Ave, Ocala, FL 34481",
    "3032 Bridgefield Ct, The Villages, FL 32162",
    "132 White Blossom Trl, Port Saint Joe, FL 32456",
    "1285 Azalea Dr, Jacksonville, FL 32205",
    "912 Anthony St E, Lehigh Acres, FL 33974",
    "1949 Par Dr, Naples, FL 34120",
    "1409 Courtside Dr, Panama City Beach, FL 32413",
    "1408 Courtside Dr, Panama City Beach, FL 32413",
    "1410 Courtside Dr, Panama City Beach, FL 32413",
    "125 Lightning Bug Ln, Sopchoppy, FL 32358",
    "5711 Peach Dr, Milton, FL 32571",
    "2506 23rd St W, Lehigh Acres, FL 33971",
    "960 NE Juniper Pl, Jensen Beach, FL 34957",
    "3408 69th St W, Lehigh Acres, FL 33971",
    "3419 26th St SW, Lehigh Acres, FL 33976",
    "3402 72nd St W, Lehigh Acres, FL 33971",
    "3506 67th St W, Lehigh Acres, FL 33971",
    "2608 45th St SW, Lehigh Acres, FL 33976",
    "15738 Switch Cane St, Clermont, FL 34711",
    "24 Riverwalk Cir, Freeport, FL 32439",
    "72 Riverwalk Cir, Freeport, FL 32439",
    "2596 Surfside Blvd, Cape Coral, FL 33914",
    "6837 Misty View Dr, Jacksonville, FL 32210",
    "1336 NW 13th Pl, Cape Coral, FL 33993",
    "10827 Venice Cir, Tampa, FL 33635",
    "6132 Highwood Park Ln, Naples, FL 34110",
    "1898 SE Coronado Ln, Stuart, FL 34996",
    "3319 Clubview Dr, North Fort Myers, FL 33917",
    "3618 Elm St, Ellenton, FL 34222",
    "2832 Vintage View Loop, Lakeland, FL 33812",
    "4734 Amhurst Cir, Destin, FL 32541",
    "4346 Marine Pkwy, New Port Richey, FL 34652",
]

def test_quote_fg_vh2(page: Page) -> None:
    username = "tsicortesmarquez@geovera"
    address = random.choice(FL_ADDRESSES)
    coverage_a = "600,000"

    page.goto("https://tst.mygeosource.com/")
    page.get_by_role("textbox", name="Policyholder Email or Agent").click()
    page.get_by_role("textbox", name="Policyholder Email or Agent").fill(username)
    page.get_by_role("textbox", name="Password").click()
    page.get_by_role("textbox", name="Password").fill(os.environ["INTERNAL_PASSWORD"])
    page.get_by_role("button", name="Log In", exact=True).click()
    page.get_by_role("paragraph").filter(has_text="Start New Quote").get_by_role("link").click()
    page.locator("#optionHomeowners").get_by_text("GeoVera Specialty Insurance").click()
    with page.expect_popup() as page1_info:
        page.get_by_role("button", name="OK").click()
    page1 = page1_info.value
    page1.wait_for_selector("#ctl00_WelcomeUserControl_ddlWholesaler", timeout=30000)
    page1.locator("#ctl00_WelcomeUserControl_ddlWholesaler").select_option("APPUND")
    page1.locator("#ctl00_WelcomeUserControl_txtProducerNumber").fill("APPAGT")
    page1.locator("#ctl00_WelcomeUserControl_btnUpdateQuotingAs").click()
    page1.wait_for_selector("#ctl00_WelcomeUserControl_ucWelcomeUserLookup_dgUserList_ctl02_btnId", timeout=30000)
    page1.locator("#ctl00_WelcomeUserControl_ucWelcomeUserLookup_dgUserList_ctl02_btnId").click()
    page1.wait_for_timeout(3000)
    page1.locator("#RadAutoCompleteBoxGeo_Input").click()
    page1.locator("#RadAutoCompleteBoxGeo_Input").press_sequentially(address)
    page1.wait_for_timeout(1000)
    page1.locator("#RadAutoCompleteBoxGeo_Input").press("Enter")
    page1.get_by_role("button", name="Next -->").click()

    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlOccupancyType").select_option("Vacant")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlVacantMoreThan90Days").select_option("No")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtPurchaseDateMonth").fill("01")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtPurchaseDateDay").fill("01")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtPurchaseDateYear").fill("2020")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtExpectedOccupancyDateMonth").fill("12")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtExpectedOccupancyDateDay").fill("31")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtExpectedOccupancyDateYear").fill("2024")
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlVacancyReason").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlVacancyReason").select_option("1")

    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkProductType_0").uncheck()
    if not page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkProductType_1").is_enabled():
        pytest.skip("chkProductType_1 is disabled for this address — skipping test")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkProductType_1").check()
    page1.get_by_role("checkbox", name="HO Wind Only").uncheck()
    page1.locator("#txtRequestedHomeownersCoverageALimit").click()
    page1.locator("#txtRequestedHomeownersCoverageALimit").fill(coverage_a)
    page1.get_by_role("button", name="Next -->").click()

    page1.locator("#ddlConstructionType").select_option("FR")
    page1.locator("#ddlSidingTypeDominant").select_option("ST")
    page1.locator("#ddlRestrictedRoofShapes").select_option("False")
    if page1.locator("#ddlRoofShapeDominant").is_visible():
        page1.locator("#ddlRoofShapeDominant").select_option("2")
    page1.locator("#ddlRoofMaterialDominant").select_option("Architectural Shingles")
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtFoundationElevation").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtFoundationElevation").fill("0")
    page1.get_by_role("button", name="Next -->").click()

    page1.wait_for_timeout(3000)
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlCentralFireAlarm").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlCentralBurglarAlarm").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlGatedGuardedCommunity").select_option("False")
    if page1.locator("#ddlSwimmingPool").is_visible():
        page1.locator("#ddlSwimmingPool").select_option(label="No Swimming Pool")
    page1.get_by_role("button", name="Next -->").click()

    page1.locator("#ddlPolybutylenePipes").select_option("False")
    page1.locator("#ddlAutomaticWaterShutoffValve").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlWillRenovate").select_option("No")
    page1.get_by_role("button", name="Next -->").click()

    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlWindMitigationVerificationInspection").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlWindMitigationVerificationInspection").select_option("Yes")
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtWindMitigationVerificationInspectionMonth").fill("01")
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtWindMitigationVerificationInspectionDay").fill("01")
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtWindMitigationVerificationInspectionYear").fill("2022")
        page1.locator("#fileWindMitigationVerificationInspectionForm").set_input_files("C:/Users/icortesmarquez/Desktop/Playwright/HO_PW/uploads/testfile.pdf")
    page1.get_by_role("button", name="Next -->").click()

    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ddlPolicyOwnerType").select_option("Individual")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtFirstName").press_sequentially("ICMTEST")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtLastName").click()
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtLastName").press_sequentially("TESTER")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtDOBMonth").click()
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtDOBMonth").press_sequentially("01")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtDOBDay").press_sequentially("01")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_txtDOBYear").press_sequentially("1998")
    page1.locator("input[name=\"ctl00\\$ContentPlaceHolder1\\$ctl05\\$txtSSN1\"]").press_sequentially("123")
    page1.locator("input[name=\"ctl00\\$ContentPlaceHolder1\\$ctl05\\$txtSSN2\"]").press_sequentially("11")
    page1.locator("input[name=\"ctl00\\$ContentPlaceHolder1\\$ctl05\\$txtSSN3\"]").press_sequentially("1122")
    page1.get_by_role("checkbox", name="Is Mailing Address the same").check()
    page1.locator("#ddlPriorHomeownersInsurance").select_option("False")
    page1.wait_for_timeout(3000)
    page1.locator("#ddlPriorHomeownersInsurance").select_option("False")
    page1.get_by_role("button", name="Next -->").click()

    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl00_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl01_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl02_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl03_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl04_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl05_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl06_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl07_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl08_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl09_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl10_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl11_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl12_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl13_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl14_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl15_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl16_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl17_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl18_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl19_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl20_ddlQuestions").select_option("False")
    page1.locator("#ctl00_ContentPlaceHolder1_ctl05_ucGenInfoQuestDDL_rptrQuestions_ctl21_ddlQuestions").select_option("False")
    page1.get_by_role("cell", name="Applicant acknowledges that flood coverage is excluded for the homeowners, wind-only, and vacant policy.", exact=True).locator("span").first.click()
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkRoofPaymentSchedule").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkRoofPaymentSchedule").check()
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkFloodExclusion").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkFloodExclusion").check()
    if page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkRoofExclusion").is_visible():
        page1.locator("#ctl00_ContentPlaceHolder1_ctl05_chkRoofExclusion").check()
    page1.get_by_role("button", name="Next -->").click()

    page1.get_by_role("button", name="Submit Quote Request").click()

    page1.wait_for_selector("#ctl00_ContentPlaceHolder1_ucQuoteDetailHOFlood_lblQuote1Header", timeout=30000)
    ProductName = page1.locator("#ctl00_ContentPlaceHolder1_ucQuoteDetailHOFlood_lblQuote1Header").inner_text()
    print(f'Homeowners Product Option: {ProductName}')
    QuoteNumber = page1.locator("#ctl00_ContentPlaceHolder1_ucQuoteDetailHOFlood_lblQuote1Number").inner_text()
    print(f'Quote Number: {QuoteNumber}')
    FullPay = page1.locator("#ctl00_ContentPlaceHolder1_ucQuoteDetailHOFlood_lstPaymentOptions_ctl00_lblQuote1PaymentOption").inner_text()
    print(f'Full Pay Option: {FullPay}')
    premium_match = re.search(r"\$\d{1,3}(?:,\d{3})*(?:\.\d{2})", FullPay)
    premium = premium_match.group() if premium_match else FullPay

    ThreePayFirstInstallment = page1.locator(
        '#ctl00_ContentPlaceHolder1_ucQuoteDetailHOFlood_lstPaymentOptions_ctl01_lblQuote1PaymentOption'
    ).text_content().strip()

    match = re.search(r"\$\d{1,3}(?:,\d{3})*(?:\.\d{2})", ThreePayFirstInstallment)

    if match:
        print(f"3 Pay First Installment:\n{match.group().strip()}\n")
    else:
        print("No dollar amount found in:", ThreePayFirstInstallment)

    page1.get_by_role("button", name="Continue to Bind Coverage").click()
    page1.locator("#chkPopInConfirmation").check()
    page1.get_by_role("button", name="Continue", exact=True).click()

    eff_date = "04/30/2026"
    page1.locator("#ctl00_ContentPlaceHolder1_ucBinderRequest_txtBinderEffectiveDate1").fill(eff_date)
    page1.locator("#ctl00_ContentPlaceHolder1_ucBinderRequest_ddlBillTo1").select_option("0")
    page1.locator("#ddlPaymentPlan1").select_option("0")
    page1.get_by_role("button", name="Request to Bind").click()

    page1.wait_for_selector("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblAmountDue", timeout=30000)
    BinderNumber = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblBinderNumber").inner_text()
    print(f'Binder Number: {BinderNumber}')
    BillTo = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblBillTo").inner_text()
    print(f'Bill To: {BillTo}')
    PaymentPlan = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblPaymentPlan").inner_text()
    print(f'Payment Plan: {PaymentPlan}')
    EffectiveDate = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblEffectiveDate").inner_text()
    print(f'Effective Date: {EffectiveDate}')
    AmountDueNow = page1.locator("#ctl00_ContentPlaceHolder1_ucBinderDetail_lblAmountDue").inner_text()
    print(f'Amount Due Now: {AmountDueNow}')
    page1.wait_for_timeout(5000)

    _append_to_xlsx([username, address, ProductName, coverage_a, QuoteNumber, premium])