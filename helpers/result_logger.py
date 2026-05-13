import os
import time
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def append_to_xlsx(xlsx_path: str, row: list, headers: list | None = None) -> None:
    """Append a row to an xlsx workbook with file-level locking for parallel-safe writes.

    Raises TimeoutError if the lock cannot be acquired within ~10 seconds. Writing
    without the lock would risk corrupting the xlsx, so we fail loudly instead.
    """
    path = os.path.abspath(xlsx_path)
    lock_path = path + ".lock"
    acquired = False
    for _ in range(20):
        try:
            fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.close(fd)
            acquired = True
            break
        except FileExistsError:
            time.sleep(0.5)
    if not acquired:
        raise TimeoutError(f"Could not acquire lock on {lock_path} after 10 seconds")
    try:
        if os.path.exists(path):
            wb = load_workbook(path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            if headers:
                ws.append(headers)
                header_fill = PatternFill(fill_type="solid", fgColor="1E5631")
                header_font = Font(bold=True, color="FFFFFF")
                for col_idx, header in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    ws.column_dimensions[cell.column_letter].width = max(len(str(header)), 10) + 2
                ws.freeze_panes = "A2"
        ws.append(row)
        wb.save(path)
    finally:
        try:
            os.unlink(lock_path)
        except Exception:
            pass


def update_xlsx_row(xlsx_path: str, key_col: int, key_value: str, updates: dict[int, str]) -> None:
    """Update a row in an xlsx workbook by matching a key column value.

    Args:
        xlsx_path: Path to the xlsx file
        key_col: Column index (1-based) to match on
        key_value: Value to match in key_col
        updates: Dict mapping column indices (1-based) to new values

    Raises TimeoutError if the lock cannot be acquired within ~10 seconds.
    """
    path = os.path.abspath(xlsx_path)
    lock_path = path + ".lock"
    acquired = False
    for _ in range(20):
        try:
            fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.close(fd)
            acquired = True
            break
        except FileExistsError:
            time.sleep(0.5)
    if not acquired:
        raise TimeoutError(f"Could not acquire lock on {lock_path} after 10 seconds")
    try:
        if not os.path.exists(path):
            raise FileNotFoundError(f"File {path} does not exist")
        wb = load_workbook(path)
        ws = wb.active
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            if row[key_col - 1].value == key_value:
                for col_idx, new_value in updates.items():
                    row[col_idx - 1].value = new_value
                wb.save(path)
                return
        raise ValueError(f"No row found with {key_col}={key_value}")
    finally:
        try:
            os.unlink(lock_path)
        except Exception:
            pass


def make_xlsx_appender(path: str, headers: list):
    """Return a parallel-safe row-appender bound to a specific xlsx file and header schema."""
    def appender(row: list) -> None:
        append_to_xlsx(path, row, headers)
    return appender


def make_xlsx_updater(path: str, key_col: int):
    """Return a parallel-safe row-updater bound to a specific xlsx file and key column.

    Args:
        path: Path to xlsx file
        key_col: Column index (1-based) to match on

    Returns a function that takes (key_value, updates) where updates is a dict of {col_idx: value}
    """
    def updater(key_value: str, updates: dict[int, str]) -> None:
        update_xlsx_row(path, key_col, key_value, updates)
    return updater
